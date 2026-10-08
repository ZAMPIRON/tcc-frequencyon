"""app.py - FrequencyON: create_app, rotas e integração do totem com o banco."""
import csv
import io
import os
import uuid
from datetime import date, timedelta
from functools import wraps

import cv2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from flask import (Flask, Blueprint, render_template, request, redirect, url_for,
                   flash, session, jsonify, Response)
from sqlalchemy import or_

from email_service import carregar_env, enviar_justificativa

carregar_env(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

from database import (db, Admin, Professor, Turma, Aluno, Presenca, ProfessorTurma, Justificativa,
                      AULAS_POR_DIA, DIAS_SELECIONAVEIS, NOMES_DIAS,
                      hoje_local, agora_local, eh_dia_com_aula, ultimo_dia_com_aula,
                      dias_de_formulario, migrar_professor_turma)
from face_service import Totem, encoding_de_arquivo

admin = Blueprint("admin", __name__, url_prefix="/admin")
professor = Blueprint("professor", __name__, url_prefix="/professor")
area_aluno = Blueprint("aluno", __name__, url_prefix="/aluno")
reconhecimento = Blueprint("reconhecimento", __name__)
totem = None  # definido em create_app


# ============================ helpers ============================
def salvar_foto(arquivo):
    """Salva upload em static/uploads e devolve a URL (ou None)."""
    if not arquivo or not arquivo.filename:
        return None
    ext = os.path.splitext(arquivo.filename)[1].lower() or ".jpg"
    nome = f"{uuid.uuid4().hex}{ext}"
    arquivo.save(os.path.join(UPLOAD_DIR(), nome))
    return f"/static/uploads/{nome}"


def UPLOAD_DIR():
    from flask import current_app
    return current_app.config["UPLOAD_DIR"]


def encoding_de_foto(url):
    """Gera o encoding facial a partir de uma foto já salva. None se não achar 1 rosto."""
    if not url:
        return None
    try:
        enc = encoding_de_arquivo(os.path.join(UPLOAD_DIR(), os.path.basename(url)))
        return enc.tobytes() if enc is not None else None
    except Exception as e:
        print(f"[foto] não consegui ler o rosto da foto: {e}")
        return None


def usuario_atual():
    tipo, uid = session.get("tipo"), session.get("uid")
    modelo = {"admin": Admin, "professor": Professor, "aluno": Aluno}.get(tipo)
    return db.session.get(modelo, uid) if modelo and uid else None


def requer(tipo):
    def deco(f):
        @wraps(f)
        def wrapper(*a, **kw):
            if session.get("tipo") != tipo or not usuario_atual():
                return redirect(url_for({"admin": "login_admin", "professor": "login_professor"}
                                        .get(tipo, "login_aluno")))
            return f(*a, **kw)
        return wrapper
    return deco


# ============================ TOTEM ============================
@reconhecimento.route("/totem")
def pagina_totem():
    totem.iniciar()
    return render_template("totem/reconhecimento.html")


@reconhecimento.route("/video_feed")
def video_feed():
    totem.iniciar()
    return Response(totem.stream(), mimetype="multipart/x-mixed-replace; boundary=frame")


@reconhecimento.route("/api/reconhecimento")
def api_reconhecimento():
    ev = totem.proximo_evento()   # cada evento é entregue uma única vez
    return jsonify({"reconhecido": True, "dados": ev} if ev else {"reconhecido": False})


@reconhecimento.route("/api/cadastrar_foto", methods=["POST"])
def api_cadastrar_foto():
    try:
        matricula = ((request.get_json(silent=True) or {}).get("matricula") or "").strip()
        aluno = Aluno.query.filter_by(matricula=matricula).first()
        if not aluno:
            return jsonify(sucesso=False, mensagem="Matrícula não encontrada."), 404
        totem.iniciar()
        enc, frame = totem.capturar_encoding()
        if enc is None:
            return jsonify(sucesso=False, mensagem=frame), 422
        aluno.face_encoding = enc.tobytes()
        if not aluno.foto:  # usa a captura como foto de perfil
            nome = f"{uuid.uuid4().hex}.jpg"
            cv2.imwrite(os.path.join(UPLOAD_DIR(), nome), frame)
            aluno.foto = f"/static/uploads/{nome}"
        db.session.commit()
        totem.recarregar()
        return jsonify(sucesso=True, mensagem=f"Rosto de {aluno.nome} cadastrado com sucesso!")
    except Exception as e:
        db.session.rollback()
        print(f"[cadastrar_foto] erro: {e}")
        return jsonify(sucesso=False, mensagem="Erro interno ao cadastrar o rosto."), 500


# ============================ ADMIN ============================
@admin.route("/")
@admin.route("/dashboard")
@requer("admin")
def dashboard():
    turmas, alunos = Turma.query.all(), Aluno.query.all()
    por_curso = {}
    for t in turmas:
        por_curso.setdefault(t.curso, []).append(t.frequencia_media)
    freqs = [a.frequencia for a in alunos]
    return render_template(
        "admin/dashboard.html",
        total_alunos=len(alunos), total_professores=Professor.query.count(),
        total_turmas=len(turmas),
        frequencia_geral=round(sum(freqs) / len(freqs)) if freqs else 0,
        frequencia_por_curso=[(c, sum(v) / len(v)) for c, v in por_curso.items()],
        turmas_criticas=sorted(turmas, key=lambda t: t.frequencia_media)[:5],
        alunos_em_risco=sum(1 for x in alunos if x.presencas and x.frequencia < 75))


@admin.route("/alunos")
@requer("admin")
def alunos():
    busca = request.args.get("busca", "").strip()
    q = Aluno.query
    if busca:
        q = q.filter(or_(Aluno.nome.ilike(f"%{busca}%"), Aluno.matricula.ilike(f"%{busca}%")))
    return render_template("admin/alunos.html", alunos=q.order_by(Aluno.nome).all(), busca=busca)


@admin.route("/alunos/novo", methods=["GET", "POST"])
@requer("admin")
def novo_aluno():
    if request.method == "POST":
        f = request.form
        if Aluno.query.filter_by(matricula=f["matricula"].strip()).first():
            flash("Já existe um aluno com essa matrícula.", "danger")
            return redirect(url_for("admin.novo_aluno"))
        a = Aluno(nome=f["nome"].strip(), email=f.get("email"), matricula=f["matricula"].strip(),
                  numero_chamada=int(f["numero_chamada"]) if f.get("numero_chamada") else None,
                  turma_id=int(f["turma_id"]) if f.get("turma_id") else None)
        a.set_senha(f["senha"])
        a.foto = salvar_foto(request.files.get("foto"))
        a.face_encoding = encoding_de_foto(a.foto)     # já deixa o rosto pronto p/ o totem
        db.session.add(a)
        db.session.commit()
        totem.recarregar()
        if a.foto and not a.face_encoding:
            flash("Aluno cadastrado, mas não achei um rosto único na foto. "
                  "Cadastre o rosto pelo totem.", "warning")
        else:
            flash("Aluno cadastrado com sucesso!", "success")
        return redirect(url_for("admin.alunos"))
    return render_template("admin/cadastro_aluno.html", turmas=Turma.query.all())


@admin.route("/alunos/<int:id>/deletar", methods=["POST"])
@requer("admin")
def deletar_aluno(id):
    db.session.delete(db.get_or_404(Aluno, id))
    db.session.commit()
    totem.recarregar()
    flash("Aluno removido.", "success")
    return redirect(url_for("admin.alunos"))


@admin.route("/professores", methods=["GET", "POST"])
@requer("admin")
def professores():
    if request.method == "POST":
        f = request.form
        if Professor.query.filter_by(email=f["email"]).first():
            flash("Já existe um professor com esse e-mail.", "danger")
        else:
            p = Professor(nome=f["nome"], email=f["email"], departamento=f.get("departamento"),
                          foto=salvar_foto(request.files.get("foto")))
            p.set_senha(f["senha"])
            db.session.add(p)
            db.session.commit()
            flash("Professor cadastrado!", "success")
        return redirect(url_for("admin.professores"))
    return render_template("admin/professores.html", professores=Professor.query.order_by(Professor.nome).all())


@admin.route("/professores/<int:id>/deletar", methods=["POST"])
@requer("admin")
def deletar_professor(id):
    p = db.get_or_404(Professor, id)
    db.session.delete(p)   # os vínculos com turmas são apagados junto
    db.session.commit()
    return redirect(url_for("admin.professores"))


@admin.route("/turmas", methods=["GET", "POST"])
@requer("admin")
def turmas():
    if request.method == "POST":
        f = request.form
        turma = Turma(nome=f["nome"], curso=f["curso"])
        for pid in set(f.getlist("professor_ids", type=int)):
            if db.session.get(Professor, pid):
                turma.vinculos.append(ProfessorTurma(professor_id=pid, dias=""))
        db.session.add(turma)
        db.session.commit()
        flash("Turma criada! Defina os dias de aula de cada professor nos detalhes da turma.", "success")
        return redirect(url_for("admin.turmas"))
    return render_template("admin/turmas.html", turmas=Turma.query.all(),
                           professores=Professor.query.order_by(Professor.nome).all())


@admin.route("/turmas/<int:id>")
@requer("admin")
def detalhes_turma(id):
    turma = db.get_or_404(Turma, id)
    vinculados = {v.professor_id for v in turma.vinculos}
    return render_template("admin/detalhes_turma.html", turma=turma,
                           alunos_sem_turma=Aluno.query.filter_by(turma_id=None).all(),
                           professores_disponiveis=[p for p in Professor.query.order_by(Professor.nome)
                                                    if p.id not in vinculados],
                           dias_opcoes=[(d, NOMES_DIAS[d]) for d in DIAS_SELECIONAVEIS])


@admin.route("/turmas/<int:turma_id>/professores", methods=["POST"])
@requer("admin")
def vincular_professor(turma_id):
    turma = db.get_or_404(Turma, turma_id)
    prof = db.get_or_404(Professor, request.form.get("professor_id", type=int))
    if not db.session.get(ProfessorTurma, (prof.id, turma.id)):
        v = ProfessorTurma(professor_id=prof.id, turma_id=turma.id)
        v.dias_semana = dias_de_formulario(request.form.getlist("dias"))
        db.session.add(v)
        db.session.commit()
        flash(f"{prof.nome} vinculado à turma.", "success")
    return redirect(url_for("admin.detalhes_turma", id=turma_id))


@admin.route("/turmas/<int:turma_id>/professores/<int:professor_id>/dias", methods=["POST"])
@requer("admin")
def dias_professor_turma(turma_id, professor_id):
    v = db.get_or_404(ProfessorTurma, (professor_id, turma_id))
    v.dias_semana = dias_de_formulario(request.form.getlist("dias"))
    db.session.commit()
    flash("Dias de aula atualizados.", "success")
    return redirect(url_for("admin.detalhes_turma", id=turma_id))


@admin.route("/turmas/<int:turma_id>/professores/<int:professor_id>/remover", methods=["POST"])
@requer("admin")
def desvincular_professor(turma_id, professor_id):
    db.session.delete(db.get_or_404(ProfessorTurma, (professor_id, turma_id)))
    db.session.commit()
    flash("Professor removido da turma.", "success")
    return redirect(url_for("admin.detalhes_turma", id=turma_id))


@admin.route("/turmas/<int:id>/deletar", methods=["POST"])
@requer("admin")
def deletar_turma(id):
    t = db.get_or_404(Turma, id)
    for a in t.alunos:
        a.turma_id = None
    db.session.delete(t)
    db.session.commit()
    return redirect(url_for("admin.turmas"))


@admin.route("/turmas/<int:turma_id>/adicionar", methods=["POST"])
@requer("admin")
def adicionar_aluno_turma(turma_id):
    aluno = db.get_or_404(Aluno, int(request.form["aluno_id"]))
    aluno.turma_id = turma_id
    db.session.commit()
    return redirect(url_for("admin.detalhes_turma", id=turma_id))


@admin.route("/alunos/<int:aluno_id>/desvincular", methods=["POST"])
@requer("admin")
def desvincular_aluno(aluno_id):
    aluno = db.get_or_404(Aluno, aluno_id)
    turma_id, aluno.turma_id = aluno.turma_id, None
    db.session.commit()
    return redirect(url_for("admin.detalhes_turma", id=turma_id) if turma_id else url_for("admin.turmas"))


@admin.route("/risco-evasao")
@requer("admin")
def risco_evasao():
    return render_template("admin/risco_evasao.html",
                           alunos=sorted([a for a in Aluno.query.all() if a.presencas and a.frequencia < 75],
                                         key=lambda a: a.frequencia))


def _dados_relatorio():
    """Indicadores usados na página de relatórios e na planilha exportada."""
    alunos, turmas = Aluno.query.all(), Turma.query.all()
    freqs = [a.frequencia for a in alunos]
    por_curso = {}
    for t in turmas:
        por_curso.setdefault(t.curso, []).append(t.frequencia_media)
    dias_por_turma = {t.id: t.dias_aula for t in turmas}
    meses = {}
    for p in Presenca.query.all():
        if not eh_dia_com_aula(p.data, dias_por_turma.get(p.aluno.turma_id, ())):
            continue
        ok, tot = meses.get(p.data.strftime("%Y-%m"), (0, 0))
        meses[p.data.strftime("%Y-%m")] = (ok + (p.status != "falta"), tot + 1)
    chaves = sorted(meses)[-6:]
    n = len(alunos)
    em_risco = sum(1 for x in alunos if x.presencas and x.frequencia < 75)
    return dict(
        alunos=alunos, turmas=turmas, total_alunos=n,
        media_geral=round(sum(freqs) / n) if n else 0,
        alunos_criticos=em_risco,
        taxa_evasao=round(100 * em_risco / n, 1) if n else 0,
        cursos_labels=list(por_curso), cursos_valores=[round(sum(v) / len(v)) for v in por_curso.values()],
        meses_labels=[k[5:] + "/" + k[:4] for k in chaves],
        meses_valores=[round(100 * meses[k][0] / meses[k][1]) for k in chaves],
        top_professores=sorted(Professor.query.all(), key=lambda p: p.media, reverse=True)[:5])


@admin.route("/relatorios")
@requer("admin")
def relatorios():
    return render_template("admin/relatorios.html", **_dados_relatorio())


def _situacao(freq):
    return "Regular" if freq >= 75 else ("Atenção" if freq >= 50 else "Crítico")


@admin.route("/relatorios/exportar")
@requer("admin")
def exportar_relatorio():
    """Planilha Excel com resumo, alunos, turmas, professores, alunos em risco e justificativas."""
    dados = _dados_relatorio()
    wb = Workbook()

    titulo_fonte = Font(bold=True, size=16, color="3B2FBF")
    cab_fonte = Font(bold=True, color="FFFFFF")
    cab_fundo = PatternFill("solid", fgColor="5B4BFF")
    secao_fonte = Font(bold=True, size=12, color="3B2FBF")
    cores_situacao = {"Regular": "D1FAE5", "Atenção": "FEF3C7", "Crítico": "FEE2E2"}
    cores_justificativa = {"aceita": "D1FAE5", "pendente": "FEF3C7", "recusada": "FEE2E2"}

    def tabela(ws, linha, cabecalho, linhas, larguras, formatos=None, coluna_cor=None, cores=None):
        """Escreve uma tabela a partir de `linha`; formatos = {índice da coluna: formato numérico}."""
        for c, texto in enumerate(cabecalho, 1):
            cel = ws.cell(row=linha, column=c, value=texto)
            cel.font, cel.fill = cab_fonte, cab_fundo
            cel.alignment = Alignment(vertical="center", wrap_text=True)
        for i, valores in enumerate(linhas, linha + 1):
            for c, v in enumerate(valores, 1):
                cel = ws.cell(row=i, column=c, value=v)
                if formatos and (c - 1) in formatos:
                    cel.number_format = formatos[c - 1]
            if coluna_cor is not None and valores[coluna_cor] in cores:
                fundo = PatternFill("solid", fgColor=cores[valores[coluna_cor]])
                for c in range(1, len(valores) + 1):
                    ws.cell(row=i, column=c).fill = fundo
        for c, largura in enumerate(larguras, 1):
            ws.column_dimensions[get_column_letter(c)].width = max(
                ws.column_dimensions[get_column_letter(c)].width or 0, largura)
        return linha + len(linhas) + 1

    def contagens(aluno):
        dias = aluno.turma.dias_aula if aluno.turma else ()
        st = [p.status for p in aluno.presencas if eh_dia_com_aula(p.data, dias)]
        return st.count("presente"), st.count("atraso"), st.count("falta")

    alunos = sorted(dados["alunos"], key=lambda a: ((a.turma.nome if a.turma else "~"), a.nome))
    linhas_alunos = []
    for a in alunos:
        pres, atr, fal = contagens(a)
        linhas_alunos.append([a.matricula, a.nome, a.turma.nome if a.turma else "Sem turma",
                              a.turma.curso if a.turma else "-", a.turma.dias_aula_nomes if a.turma else "-",
                              pres, atr, fal, a.frequencia / 100,
                              _situacao(a.frequencia) if a.presencas else "Sem registros"])

    # ---------- Resumo ----------
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "FrequencyON - Relatório de Frequência"
    ws["A1"].font = titulo_fonte
    ws["A2"] = "Gerado em " + agora_local().strftime("%d/%m/%Y %H:%M")
    ws["A2"].font = Font(italic=True, color="6B7280")

    ws["A4"] = "Indicadores gerais"
    ws["A4"].font = secao_fonte
    indicadores = [
        ["Alunos matriculados", dados["total_alunos"]],
        ["Turmas", len(dados["turmas"])],
        ["Professores", Professor.query.count()],
        ["Frequência média da instituição", dados["media_geral"] / 100],
        ["Alunos em risco (frequência < 75%)", dados["alunos_criticos"]],
        ["Taxa de risco", dados["taxa_evasao"] / 100],
        ["Justificativas pendentes", Justificativa.query.filter_by(status="pendente").count()],
    ]
    linha = tabela(ws, 5, ["Indicador", "Valor"], indicadores, [40, 16])
    ws.cell(row=9, column=2).number_format = "0%"
    ws.cell(row=11, column=2).number_format = "0.0%"

    linha += 1
    ws.cell(row=linha, column=1, value="Frequência mensal (últimos 6 meses com aula)").font = secao_fonte
    linha = tabela(ws, linha + 1, ["Mês", "Frequência"],
                   [[m, v / 100] for m, v in zip(dados["meses_labels"], dados["meses_valores"])],
                   [40, 16], {1: "0%"})
    linha += 1
    ws.cell(row=linha, column=1, value="Frequência por curso").font = secao_fonte
    tabela(ws, linha + 1, ["Curso", "Frequência"],
           [[c, v / 100] for c, v in zip(dados["cursos_labels"], dados["cursos_valores"])],
           [40, 16], {1: "0%"})

    # ---------- Alunos ----------
    cab_alunos = ["Matrícula", "Nome", "Turma", "Curso", "Dias de aula",
                  "Aulas presente", "Aulas com atraso", "Aulas com falta", "Frequência", "Situação"]
    larg_alunos = [12, 32, 14, 34, 18, 14, 16, 15, 12, 14]
    ws = wb.create_sheet("Alunos")
    tabela(ws, 1, cab_alunos, linhas_alunos, larg_alunos, {8: "0%"}, 9, cores_situacao)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---------- Turmas ----------
    ws = wb.create_sheet("Turmas")
    tabela(ws, 1, ["Turma", "Curso", "Professores", "Dias de aula", "Alunos", "Frequência média", "Alunos em risco"],
           [[t.nome, t.curso, t.nomes_professores, t.dias_aula_nomes, len(t.alunos), t.frequencia_media / 100,
             sum(1 for a in t.alunos if a.presencas and a.frequencia < 75)]
            for t in sorted(dados["turmas"], key=lambda t: t.nome)],
           [14, 34, 30, 18, 9, 17, 15], {5: "0.0%"})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---------- Professores ----------
    ws = wb.create_sheet("Professores")
    tabela(ws, 1, ["Professor", "E-mail", "Departamento", "Turmas (dias de aula)", "Frequência média das turmas"],
           [[p.nome, p.email, p.departamento or "-",
             "; ".join(f"{v.turma.nome} ({v.dias_nomes})"
                       for v in sorted(p.vinculos, key=lambda v: v.turma.nome)) or "-",
             p.media / 100]
            for p in sorted(Professor.query.all(), key=lambda p: p.nome)],
           [28, 30, 18, 40, 16], {4: "0%"})
    ws.freeze_panes = "A2"

    # ---------- Em risco ----------
    ws = wb.create_sheet("Em risco")
    risco = sorted((l for l in linhas_alunos if l[9] in ("Atenção", "Crítico")), key=lambda l: l[8])
    if risco:
        tabela(ws, 1, cab_alunos, risco, larg_alunos, {8: "0%"}, 9, cores_situacao)
        ws.freeze_panes = "A2"
    else:
        ws["A1"] = "Nenhum aluno com frequência abaixo de 75%."

    # ---------- Justificativas ----------
    ws = wb.create_sheet("Justificativas")
    js = Justificativa.query.order_by(Justificativa.data.desc()).all()
    if js:
        tabela(ws, 1, ["Data da falta", "Aluno", "Matrícula", "Turma", "Motivo", "Situação",
                       "Enviada em", "Respondida em"],
               [[j.data, j.aluno.nome, j.aluno.matricula, j.aluno.turma.nome if j.aluno.turma else "-",
                 j.motivo, j.status, j.criada_em, j.respondida_em] for j in js],
               [13, 28, 12, 12, 50, 12, 17, 17],
               {0: "DD/MM/YYYY", 6: "DD/MM/YYYY HH:MM", 7: "DD/MM/YYYY HH:MM"},
               5, cores_justificativa)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    else:
        ws["A1"] = "Nenhuma justificativa enviada."

    saida = io.BytesIO()
    wb.save(saida)
    nome = f"relatorio_frequencyon_{hoje_local().isoformat()}.xlsx"
    return Response(saida.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={nome}"})


# ============================ PROFESSOR ============================
def _normalizar_data_chamada(valor, dias):
    """Retorna uma data válida de aula: nunca futura e sempre em dia de aula."""
    hoje = hoje_local()

    if not valor:
        selecionada = hoje
    else:
        try:
            selecionada = date.fromisoformat(valor)
        except (TypeError, ValueError):
            selecionada = hoje

    if selecionada > hoje:
        selecionada = hoje

    # dias sem aula apontam para o dia de aula anterior mais próximo
    return ultimo_dia_com_aula(selecionada, dias)


MESES = ("Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro")


def _fim_do_mes(d):
    return (d.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)


def _link_mes(d, hoje, dias, turma_id):
    """Link para outro mês: abre no último dia de aula daquele mês (sem passar de hoje)."""
    alvo = ultimo_dia_com_aula(min(_fim_do_mes(d), hoje), dias)
    return url_for("professor.chamada", turma_id=turma_id, data=alvo.isoformat())


def _calendario_mes(data_ref, hoje, dias_aula, turma, taxa_por_dia):
    """Monta o calendário do mês de data_ref (semanas começando na segunda)."""
    primeiro = data_ref.replace(day=1)
    ultimo = _fim_do_mes(data_ref)
    d = primeiro - timedelta(days=primeiro.weekday())
    semanas = []
    while d <= ultimo:
        semana = []
        for _ in range(7):
            aula = d.month == data_ref.month and eh_dia_com_aula(d, dias_aula)
            taxa = taxa_por_dia.get(d)
            nivel = None
            if aula and d <= hoje:
                nivel = "vazio" if taxa is None else ("bom" if taxa >= 75 else "medio" if taxa >= 50 else "ruim")
            semana.append({
                "dia": d.day,
                "no_mes": d.month == data_ref.month,
                "aula": aula,
                "futuro": d > hoje,
                "hoje": d == hoje,
                "selecionado": d == data_ref,
                "nivel": nivel,
                "taxa": taxa,
                "url": url_for("professor.chamada", turma_id=turma.id, data=d.isoformat())
                       if aula and d <= hoje else None,
            })
            d += timedelta(days=1)
        semanas.append(semana)

    mes_anterior = primeiro - timedelta(days=1)
    proximo_mes = ultimo + timedelta(days=1)
    return {
        "titulo": f"{MESES[data_ref.month - 1]} {data_ref.year}",
        "semanas": semanas,
        "cabecalho": [n[:3] for n in NOMES_DIAS],
        "anterior": _link_mes(mes_anterior, hoje, dias_aula, turma.id),
        "proximo": _link_mes(proximo_mes, hoje, dias_aula, turma.id) if proximo_mes <= hoje else None,
        "hoje": url_for("professor.chamada", turma_id=turma.id),
    }


@professor.route("/chamada")
@requer("professor")
def chamada():
    prof = usuario_atual()
    turma_id = request.args.get("turma_id", type=int)
    hoje = hoje_local()

    turmas = prof.turmas
    turma = next((t for t in turmas if t.id == turma_id), turmas[0] if turmas else None)
    vinculo = prof.vinculo(turma.id) if turma else None
    dias_prof = vinculo.dias_semana if vinculo else ()
    data_selecionada = _normalizar_data_chamada(request.args.get("data"), dias_prof)
    alunos, presentes, faltas = [], 0, 0
    calendario = None

    if turma and dias_prof:
        ids = [a.id for a in turma.alunos]
        inicio_mes, fim_mes = data_selecionada.replace(day=1), _fim_do_mes(data_selecionada)
        registros = Presenca.query.filter(
            Presenca.aluno_id.in_(ids),
            Presenca.data.between(inicio_mes, fim_mes)
        ).all() if ids else []

        # chamada do dia selecionado + totais do mês (só dias de aula do professor)
        mapa, ok_mes, ok_dia, dias_com_chamada = {}, {}, {}, set()
        for r in registros:
            if not eh_dia_com_aula(r.data, dias_prof):
                continue
            dias_com_chamada.add(r.data)
            presente = r.status in ("presente", "atraso")
            ok_mes[r.aluno_id] = ok_mes.get(r.aluno_id, 0) + presente
            ok_dia[r.data] = ok_dia.get(r.data, 0) + presente
            if r.data == data_selecionada:
                mapa.setdefault(r.aluno_id, {})[r.aula] = r.status

        aulas_por_dia_turma = len(ids) * AULAS_POR_DIA
        taxa_por_dia = {d: round(100 * ok_dia.get(d, 0) / aulas_por_dia_turma)
                        for d in dias_com_chamada}
        calendario = _calendario_mes(data_selecionada, hoje, dias_prof, turma, taxa_por_dia)
        total_aulas_mes = len(dias_com_chamada) * AULAS_POR_DIA

        for a in sorted(turma.alunos, key=lambda x: (x.numero_chamada or 999, x.nome)):
            st = [mapa.get(a.id, {}).get(i, "falta") for i in range(AULAS_POR_DIA)]
            tem_presenca = any(s in ("presente", "atraso") for s in st)
            faltou_dia = all(s == "falta" for s in st)
            presentes += int(tem_presenca)
            faltas += int(faltou_dia)
            alunos.append({
                "id": a.id,
                "nome": a.nome,
                "matricula": a.matricula,
                "iniciais": "".join(p[0] for p in a.nome.split()[:2]).upper(),
                "status": st,
                "freq_mes": round(100 * ok_mes.get(a.id, 0) / total_aulas_mes) if total_aulas_mes else None,
            })

    eh_hoje = data_selecionada == hoje
    pode_editar = eh_hoje and eh_dia_com_aula(hoje, dias_prof)

    return render_template(
        "professor/frequencia.html",
        turmas=turmas,
        turma_selecionada=turma,
        alunos=alunos,
        total_matriculados=len(alunos),
        total_presentes=presentes,
        total_faltas=faltas,
        media_frequencia=turma.frequencia_media if turma else 0,
        data_iso=data_selecionada.isoformat(),
        data_exibicao=data_selecionada.strftime("%d/%m/%Y"),
        eh_hoje=eh_hoje,
        pode_editar=pode_editar,
        dia_nome=NOMES_DIAS[data_selecionada.weekday()],
        mes_nome=MESES[data_selecionada.month - 1],
        dias_prof=dias_prof,
        dias_prof_nomes=vinculo.dias_nomes if vinculo else "",
        calendario=calendario,
    )


def _dados_postagem():
    dados = request.get_json(silent=True) or {}
    valor = request.args.get("data") or dados.get("data") or ""
    senha_admin = (dados.get("admin_senha") or "").strip()

    try:
        data_requisitada = date.fromisoformat(valor)
    except (TypeError, ValueError):
        data_requisitada = None

    return data_requisitada, senha_admin


def _senha_admin_valida(senha):
    if not senha:
        return False
    return any(admin.check_senha(senha) for admin in Admin.query.all())


def _pode_alterar_data(data_requisitada, senha_admin, dias):
    hoje = hoje_local()

    if data_requisitada is None:
        return False, "Data inválida."

    if data_requisitada > hoje:
        return False, "Não é possível alterar uma chamada futura."

    if not eh_dia_com_aula(data_requisitada, dias):
        return False, "Você não tem aula nesta turma neste dia da semana."

    if data_requisitada == hoje:
        return True, None

    if not _senha_admin_valida(senha_admin):
        return False, "Para alterar uma chamada anterior, informe a senha de um administrador."

    return True, None


def _alterar(aluno_id, barra, regra):
    if barra < 0 or barra >= AULAS_POR_DIA:
        return jsonify(erro="Aula inválida."), 400

    aluno = db.get_or_404(Aluno, aluno_id)
    vinculo = db.session.get(ProfessorTurma, (session["uid"], aluno.turma_id)) if aluno.turma_id else None
    if not vinculo:
        return jsonify(erro="Aluno não pertence às suas turmas."), 403

    data_requisitada, senha_admin = _dados_postagem()
    permitido, erro = _pode_alterar_data(data_requisitada, senha_admin, vinculo.dias_semana)
    if not permitido:
        return jsonify(erro=erro), 403

    p = Presenca.query.filter_by(
        aluno_id=aluno_id, data=data_requisitada, aula=barra
    ).first()

    if not p:
        p = Presenca(
            aluno_id=aluno_id,
            data=data_requisitada,
            aula=barra,
            status="falta"
        )
        db.session.add(p)

    p.status, p.origem = regra(p.status), "professor"
    db.session.commit()

    return jsonify(estado=p.status, media_frequencia=aluno.turma.frequencia_media)


@professor.route("/toggle/<int:aluno_id>/<int:barra>", methods=["POST"])
@requer("professor")
def toggle(aluno_id, barra):
    return _alterar(aluno_id, barra, lambda s: "falta" if s == "presente" else "presente")


@professor.route("/atraso/<int:aluno_id>/<int:barra>", methods=["POST"])
@requer("professor")
def atraso(aluno_id, barra):
    if barra > 1:
        return jsonify(erro="Atraso só pode ser marcado nas aulas 1 e 2."), 400
    return _alterar(aluno_id, barra, lambda s: "presente" if s == "atraso" else "atraso")


# ============================ ALUNO ============================
def _historico_aluno(aluno):
    """Dados da tela de detalhes: histórico por dia de aula, totais e justificativas."""
    hoje = hoje_local()
    turma = aluno.turma
    dias_aula = turma.dias_aula if turma else ()
    linhas, justificaveis = [], []
    presentes = atrasos = total = 0

    if dias_aula:
        # dias em que houve chamada na turma (mesmo critério de Aluno.frequencia)
        datas = {d for (d,) in db.session.query(Presenca.data).join(Aluno)
                 .filter(Aluno.turma_id == turma.id).distinct()
                 if eh_dia_com_aula(d, dias_aula)}
        mapa = {}
        for p in aluno.presencas:
            mapa.setdefault(p.data, {})[p.aula] = p.status
        justificativas = {j.data: j for j in aluno.justificativas}

        pendente_hoje = eh_dia_com_aula(hoje, dias_aula) and hoje not in datas
        for d in sorted(datas | ({hoje} if pendente_hoje else set()), reverse=True):
            sem_chamada = d not in datas
            st = (["pendente"] * AULAS_POR_DIA if sem_chamada else
                  [mapa.get(d, {}).get(i, "falta") for i in range(AULAS_POR_DIA)])
            if not sem_chamada:
                total += AULAS_POR_DIA
                presentes += st.count("presente")
                atrasos += st.count("atraso")
            faltas_dia = st.count("falta")
            linhas.append({
                "data": d,
                "dia_nome": NOMES_DIAS[d.weekday()],
                "professores": [v.professor for v in turma.vinculos if d.weekday() in v.dias_semana],
                "status": st,
                "faltas": faltas_dia,
                "atrasos": st.count("atraso"),
                "justificativa": justificativas.get(d),
            })
            if faltas_dia and (d not in justificativas or justificativas[d].status == "recusada"):
                justificaveis.append(d)

    pct_presenca = round(100 * presentes / total) if total else 0
    pct_atraso = round(100 * atrasos / total) if total else 0
    pct_falta = 100 - pct_presenca - pct_atraso if total else 0
    return {
        "linhas": linhas,
        "justificaveis": justificaveis,
        "total_aulas": total,
        "pct_presenca": pct_presenca,
        "pct_atraso": pct_atraso,
        "pct_falta": pct_falta,
        "frequencia": aluno.frequencia,
    }


def _render_detalhes(aluno, visao):
    return render_template("aluno/detalhes.html", aluno=aluno, visao=visao,
                           **_historico_aluno(aluno))


def _exportar_csv(aluno):
    """CSV (separado por ';', com BOM para abrir certo no Excel) com o histórico do aluno."""
    dados = _historico_aluno(aluno)
    rotulo = {"presente": "P", "atraso": "A", "falta": "F", "pendente": "-"}
    saida = io.StringIO()
    w = csv.writer(saida, delimiter=";")
    w.writerow(["Aluno", aluno.nome, "Matrícula", aluno.matricula,
                "Turma", aluno.turma.nome if aluno.turma else "-"])
    w.writerow(["Frequência", f"{dados['frequencia']}%", "Presenças", f"{dados['pct_presenca']}%",
                "Faltas", f"{dados['pct_falta']}%", "Atrasos", f"{dados['pct_atraso']}%"])
    w.writerow([])
    w.writerow(["Data", "Dia", "Professor(es)"] + [f"Aula {i + 1}" for i in range(AULAS_POR_DIA)]
               + ["Faltas", "Atrasos", "Justificativa"])
    for l in dados["linhas"]:
        j = l["justificativa"]
        w.writerow([l["data"].strftime("%d/%m/%Y"), l["dia_nome"],
                    ", ".join(p.nome for p in l["professores"])]
                   + [rotulo[s] for s in l["status"]]
                   + [l["faltas"], l["atrasos"], j.status if j else ""])
    w.writerow([])
    w.writerow(["Legenda: P = presente, A = atraso, F = falta, - = chamada ainda não lançada"])
    nome = f"frequencia_{aluno.matricula}.csv"
    return Response("﻿" + saida.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={nome}"})


@area_aluno.route("/")
@requer("aluno")
def painel():
    return _render_detalhes(usuario_atual(), "aluno")


@area_aluno.route("/exportar")
@requer("aluno")
def exportar():
    return _exportar_csv(usuario_atual())


@area_aluno.route("/justificar", methods=["POST"])
@requer("aluno")
def justificar():
    aluno = usuario_atual()
    motivo = (request.form.get("motivo") or "").strip()
    try:
        data_falta = date.fromisoformat(request.form.get("data", ""))
    except ValueError:
        data_falta = None

    if data_falta not in _historico_aluno(aluno)["justificaveis"]:
        flash("Escolha um dia com falta que ainda não foi justificado.", "danger")
    elif len(motivo) < 5:
        flash("Descreva o motivo da falta.", "danger")
    else:
        j = Justificativa.query.filter_by(aluno_id=aluno.id, data=data_falta).first()
        if not j:
            j = Justificativa(aluno_id=aluno.id, data=data_falta)
            db.session.add(j)
        j.motivo, j.status, j.criada_em, j.respondida_em = motivo, "pendente", agora_local(), None
        db.session.commit()
        enviar_justificativa(aluno, j, url_for("admin.detalhes_aluno", id=aluno.id, _external=True))
        flash("Justificativa enviada! Aguarde a análise da coordenação.", "success")
    return redirect(url_for("aluno.painel"))


@admin.route("/alunos/<int:id>")
@requer("admin")
def detalhes_aluno(id):
    return _render_detalhes(db.get_or_404(Aluno, id), "admin")


@admin.route("/alunos/<int:id>/exportar")
@requer("admin")
def exportar_aluno(id):
    return _exportar_csv(db.get_or_404(Aluno, id))


@admin.route("/justificativas/<int:id>/<acao>", methods=["POST"])
@requer("admin")
def responder_justificativa(id, acao):
    j = db.get_or_404(Justificativa, id)
    if acao in ("aceita", "recusada"):
        j.status, j.respondida_em = acao, agora_local()
        db.session.commit()
        flash(f"Justificativa {acao}.", "success")
    return redirect(url_for("admin.detalhes_aluno", id=j.aluno_id))


# ============================ APP ============================
def create_app():
    global totem
    app = Flask(__name__)
    app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "troque-esta-chave")
    app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", "sqlite:///frequencyon.db")
    app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"connect_args": {"timeout": 15}}
    app.config["UPLOAD_DIR"] = os.path.join(app.static_folder, "uploads")
    os.makedirs(app.config["UPLOAD_DIR"], exist_ok=True)
    db.init_app(app)

    totem = Totem(app, camera=int(os.getenv("TOTEM_CAMERA", "0")))
    app.register_blueprint(admin)
    app.register_blueprint(professor)
    app.register_blueprint(area_aluno)
    app.register_blueprint(reconhecimento)
    app.jinja_env.filters["inicial"] = lambda n: (n or "A")[0].upper()
    app.context_processor(lambda: {"usuario": usuario_atual()})

    def entrar(modelo, tipo, destino, template):
        if request.method == "POST":
            ident, senha = request.form["email"].strip(), request.form["senha"]
            u = modelo.query.filter(or_(modelo.email == ident,
                                        *( [Aluno.matricula == ident] if modelo is Aluno else []))).first()
            if u and u.check_senha(senha):
                session.clear()
                session.update(tipo=tipo, uid=u.id)
                return redirect(destino)
            flash("E-mail/matrícula ou senha inválidos.", "danger")
        return render_template(template)

    @app.route("/")
    def index():
        return render_template("index.html")

    @app.route("/login", methods=["GET", "POST"])
    @app.route("/login/aluno", methods=["GET", "POST"])
    def login_aluno():
        return entrar(Aluno, "aluno", url_for("aluno.painel"), "login.html")

    @app.route("/login/professor", methods=["GET", "POST"])
    def login_professor():
        return entrar(Professor, "professor", url_for("professor.chamada"), "professor/login_professor.html")

    @app.route("/login/admin", methods=["GET", "POST"])
    def login_admin():
        return entrar(Admin, "admin", url_for("admin.dashboard"), "admin/login_admin.html")

    @app.route("/logout")
    def logout():
        session.clear()
        return redirect(url_for("index"))

    with app.app_context():
        db.create_all()
        migrar_professor_turma()
        if not Admin.query.first():   # admin padrão para o primeiro acesso
            a = Admin(nome="Administrador", email="admin@frequencyon.com")
            a.set_senha("admin123")
            db.session.add(a)
            db.session.commit()
            print("Admin criado: admin@frequencyon.com / admin123")
    return app


if __name__ == "__main__":
    # debug=False: o reloader abriria a câmera duas vezes
    create_app().run(host="0.0.0.0", port=5000, threaded=True, debug=False)