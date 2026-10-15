# relatorios da coordenacao: a pagina com os graficos e a planilha do excel
# as rotas entram no mesmo blueprint do admin, entao continuam em /admin/relatorios
import io

from flask import Response, render_template
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import (AULAS_POR_DIA, Aluno, Justificativa, Presenca, Professor, Turma,
                      agora_local, eh_dia_com_aula, hoje_local)
from .admin import admin
from .comum import requer


def _situacao(freq):
    return "Regular" if freq >= 75 else ("Atenção" if freq >= 50 else "Crítico")


def _dados_relatorio():
    # numeros que aparecem na pagina de relatorios e tambem na planilha
    alunos, turmas = Aluno.query.all(), Turma.query.all()
    # aluno sem turma não tem frequência (seria 0%), então fica fora da média
    freqs = [a.frequencia for a in alunos if a.turma_id]

    por_curso = {}
    for t in turmas:
        por_curso.setdefault(t.curso, []).append(t.frequencia_media)

    # frequencia mes a mes com o mesmo criterio do Aluno.frequencia: cada dia com chamada
    # na turma vale 10 aulas para cada aluno dela, e aula sem registro conta como falta
    turma_do_aluno = {x.id: x.turma_id for x in alunos}
    dias_aula = {t.id: t.dias_aula for t in turmas}
    alunos_na_turma = {t.id: len(t.alunos) for t in turmas}
    presentes_no_dia, dias_com_chamada = {}, set()
    for p in Presenca.query.all():
        tid = turma_do_aluno.get(p.aluno_id)
        if tid is None or not eh_dia_com_aula(p.data, dias_aula.get(tid, ())):
            continue
        dias_com_chamada.add((tid, p.data))
        if p.status in ("presente", "atraso"):
            presentes_no_dia[(tid, p.data)] = presentes_no_dia.get((tid, p.data), 0) + 1
    meses = {}
    for tid, dia in dias_com_chamada:
        ok, tot = meses.get(dia.strftime("%Y-%m"), (0, 0))
        meses[dia.strftime("%Y-%m")] = (ok + presentes_no_dia.get((tid, dia), 0),
                                         tot + alunos_na_turma[tid] * AULAS_POR_DIA)
    chaves = sorted(meses)[-6:]  # so os ultimos 6 meses

    n = len(alunos)
    em_risco = sum(1 for x in alunos if x.presencas and x.frequencia < 75)
    return dict(
        alunos=alunos, turmas=turmas, total_alunos=n,
        media_geral=round(sum(freqs) / len(freqs)) if freqs else 0,
        alunos_criticos=em_risco,
        taxa_evasao=round(100 * em_risco / n, 1) if n else 0,
        cursos_labels=list(por_curso), cursos_valores=[round(sum(v) / len(v)) for v in por_curso.values()],
        meses_labels=[k[5:] + "/" + k[:4] for k in chaves],
        meses_valores=[round(100 * meses[k][0] / meses[k][1]) for k in chaves],
        top_professores=sorted(Professor.query.all(), key=lambda p: p.media, reverse=True)[:5])


@admin.route("/relatorios")
@requer("admin")
def relatorios():
    return render_template("admin/relatorios.html", **_dados_relatorio())


# ---------- planilha ----------

TITULO_FONTE = Font(bold=True, size=16, color="3B2FBF")
CAB_FONTE = Font(bold=True, color="FFFFFF")
CAB_FUNDO = PatternFill("solid", fgColor="5B4BFF")
SECAO_FONTE = Font(bold=True, size=12, color="3B2FBF")
CORES_SITUACAO = {"Regular": "D1FAE5", "Atenção": "FEF3C7", "Crítico": "FEE2E2"}
CORES_JUSTIFICATIVA = {"aceita": "D1FAE5", "pendente": "FEF3C7", "recusada": "FEE2E2"}


def _tabela(ws, linha, cabecalho, linhas, larguras, formatos=None, coluna_cor=None, cores=None):
    # escreve uma tabela na aba a partir da linha passada e devolve a proxima linha livre
    # formatos = {indice da coluna: formato do numero}, ex: {8: "0%"}
    # coluna_cor + cores pinta a linha inteira de acordo com o valor daquela coluna
    for c, texto in enumerate(cabecalho, 1):
        cel = ws.cell(row=linha, column=c, value=texto)
        cel.font, cel.fill = CAB_FONTE, CAB_FUNDO
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    for i, valores in enumerate(linhas, linha + 1):
        for c, v in enumerate(valores, 1):
            cel = ws.cell(row=i, column=c, value=v)
            if formatos and (c - 1) in formatos:
                cel.number_format = formatos[c - 1]
        if coluna_cor is not None and valores[coluna_cor] in cores:
            fundo = PatternFill("solid", fgColor=cores[valores[coluna_cor]])
            for c in range(1, len(valores) + 1):
                ws.cell(row=i, column=c).fill = fundo
    for c, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(
            ws.column_dimensions[get_column_letter(c)].width or 0, largura)
    return linha + len(linhas) + 1


def _contagens(aluno):
    # quantas aulas o aluno teve de presente, atraso e falta (so em dias de aula)
    dias = aluno.turma.dias_aula if aluno.turma else ()
    st = [p.status for p in aluno.presencas if eh_dia_com_aula(p.data, dias)]
    return st.count("presente"), st.count("atraso"), st.count("falta")


@admin.route("/relatorios/exportar")
@requer("admin")
def exportar_relatorio():
    # planilha com as abas: resumo, alunos, turmas, professores, em risco e justificativas
    dados = _dados_relatorio()
    wb = Workbook()

    alunos = sorted(dados["alunos"], key=lambda a: ((a.turma.nome if a.turma else "~"), a.nome))
    linhas_alunos = []
    for a in alunos:
        pres, atr, fal = _contagens(a)
        linhas_alunos.append([a.matricula, a.nome, a.turma.nome if a.turma else "Sem turma",
                              a.turma.curso if a.turma else "-", a.turma.dias_aula_nomes if a.turma else "-",
                              pres, atr, fal, a.frequencia / 100,
                              _situacao(a.frequencia) if a.presencas else "Sem registros"])

    # aba resumo
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "FrequencyON - Relatório de Frequência"
    ws["A1"].font = TITULO_FONTE
    ws["A2"] = "Gerado em " + agora_local().strftime("%d/%m/%Y %H:%M")
    ws["A2"].font = Font(italic=True, color="6B7280")

    ws["A4"] = "Indicadores gerais"
    ws["A4"].font = SECAO_FONTE
    indicadores = [
        ["Alunos matriculados", dados["total_alunos"]],
        ["Turmas", len(dados["turmas"])],
        ["Professores", Professor.query.count()],
        ["Frequência média da instituição", dados["media_geral"] / 100],
        ["Alunos em risco (frequência < 75%)", dados["alunos_criticos"]],
        ["Taxa de risco", dados["taxa_evasao"] / 100],
        ["Justificativas pendentes", Justificativa.query.filter_by(status="pendente").count()],
    ]
    linha = _tabela(ws, 5, ["Indicador", "Valor"], indicadores, [40, 16])
    ws.cell(row=9, column=2).number_format = "0%"     # frequencia media
    ws.cell(row=11, column=2).number_format = "0.0%"  # taxa de risco

    linha += 1
    ws.cell(row=linha, column=1, value="Frequência mensal (últimos 6 meses com aula)").font = SECAO_FONTE
    linha = _tabela(ws, linha + 1, ["Mês", "Frequência"],
                    [[m, v / 100] for m, v in zip(dados["meses_labels"], dados["meses_valores"])],
                    [40, 16], {1: "0%"})
    linha += 1
    ws.cell(row=linha, column=1, value="Frequência por curso").font = SECAO_FONTE
    _tabela(ws, linha + 1, ["Curso", "Frequência"],
            [[c, v / 100] for c, v in zip(dados["cursos_labels"], dados["cursos_valores"])],
            [40, 16], {1: "0%"})

    # aba alunos
    cab_alunos = ["Matrícula", "Nome", "Turma", "Curso", "Dias de aula",
                  "Aulas presente", "Aulas com atraso", "Aulas com falta", "Frequência", "Situação"]
    larg_alunos = [12, 32, 14, 34, 18, 14, 16, 15, 12, 14]
    ws = wb.create_sheet("Alunos")
    _tabela(ws, 1, cab_alunos, linhas_alunos, larg_alunos, {8: "0%"}, 9, CORES_SITUACAO)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # aba turmas
    ws = wb.create_sheet("Turmas")
    _tabela(ws, 1, ["Turma", "Curso", "Professores", "Dias de aula", "Alunos", "Frequência média", "Alunos em risco"],
            [[t.nome, t.curso, t.nomes_professores, t.dias_aula_nomes, len(t.alunos), t.frequencia_media / 100,
              sum(1 for a in t.alunos if a.presencas and a.frequencia < 75)]
             for t in sorted(dados["turmas"], key=lambda t: t.nome)],
            [14, 34, 30, 18, 9, 17, 15], {5: "0.0%"})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # aba professores
    ws = wb.create_sheet("Professores")
    _tabela(ws, 1, ["Professor", "E-mail", "Departamento", "Turmas (dias de aula)", "Frequência média das turmas"],
            [[p.nome, p.email, p.departamento or "-",
              "; ".join(f"{v.turma.nome} ({v.dias_nomes})"
                        for v in sorted(p.vinculos, key=lambda v: v.turma.nome)) or "-",
              p.media / 100]
             for p in sorted(Professor.query.all(), key=lambda p: p.nome)],
            [28, 30, 18, 40, 16], {4: "0%"})
    ws.freeze_panes = "A2"

    # aba em risco (atencao e critico, do pior pro melhor)
    ws = wb.create_sheet("Em risco")
    risco = sorted((l for l in linhas_alunos if l[9] in ("Atenção", "Crítico")), key=lambda l: l[8])
    if risco:
        _tabela(ws, 1, cab_alunos, risco, larg_alunos, {8: "0%"}, 9, CORES_SITUACAO)
        ws.freeze_panes = "A2"
    else:
        ws["A1"] = "Nenhum aluno com frequência abaixo de 75%."

    # aba justificativas
    ws = wb.create_sheet("Justificativas")
    js = Justificativa.query.order_by(Justificativa.data.desc()).all()
    if js:
        _tabela(ws, 1, ["Data da falta", "Aluno", "Matrícula", "Turma", "Motivo", "Situação",
                        "Enviada em", "Respondida em"],
                [[j.data, j.aluno.nome, j.aluno.matricula, j.aluno.turma.nome if j.aluno.turma else "-",
                  j.motivo, j.status, j.criada_em, j.respondida_em] for j in js],
                [13, 28, 12, 12, 50, 12, 17, 17],
                {0: "DD/MM/YYYY", 6: "DD/MM/YYYY HH:MM", 7: "DD/MM/YYYY HH:MM"},
                5, CORES_JUSTIFICATIVA)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    else:
        ws["A1"] = "Nenhuma justificativa enviada."

    saida = io.BytesIO()
    wb.save(saida)
    nome = f"relatorio_frequencyon_{hoje_local().isoformat()}.xlsx"
    return Response(saida.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={nome}"})
